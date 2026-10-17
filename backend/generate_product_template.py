import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def create_product_import_template(output_path="backend/product_import_template.xlsx"):
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Sheet 1: Products (Data Entry Sheet)
    # -------------------------------------------------------------
    ws_data = wb.active
    ws_data.title = "Products"
    ws_data.views.sheetView[0].showGridLines = True
    
    headers = [
        ("SKU", "Required. E.g. LG0831, LG0831-1. Suffixes like -1 are distinct products."),
        ("Product Name", "Required. Full description or name of the item."),
        ("Category", "Optional. E.g. Power Tools, Garden Equipment, Hardware."),
        ("Brand", "Optional. E.g. BKR, JSPL, Gardena, Bosch."),
        ("Rate (Excl GST)", "Optional. Basic selling price before tax. If blank, auto-calculated from Price (Incl GST)."),
        ("Price (Incl GST)", "Optional. Price including GST. If Rate is blank, Rate = Price / (1 + GST%/100)."),
        ("GST Rate (%)", "Optional. Tax percentage (e.g. 18, 12, 5, 28, 0). Defaults to 18."),
        ("HSN Code", "Optional. 4-8 digit HSN/SAC code (e.g. 84672100)."),
        ("Unit", "Optional. Measurement unit (PCS, SET, NOS, KG, MTR). Defaults to PCS."),
        ("Barcode", "Optional. Barcode or EAN number."),
        ("Min Stock Level", "Optional. Low stock threshold for alerts. Defaults to 0."),
        ("Reorder Level", "Optional. Minimum inventory level to trigger reorder."),
        ("Initial Stock", "Optional. Opening quantity to add into Central Warehouse. Defaults to 0."),
        ("Status", "Optional. Active or Inactive. Defaults to Active.")
    ]
    
    header_titles = [h[0] for h in headers]
    ws_data.append(header_titles)
    
    sample_rows = [
        ["LG0831", "BKR Heavy Duty Thermal Fogging Machine 15L", "Garden Equipment", "BKR", 26942.86, 28290.00, 5.0, "84248200", "PCS", "8901234567890", 5, 10, 25, "Active"],
        ["LG0831-1", "BKR Heavy Duty Fogging Machine - Battery Pack Only", "Spare Parts", "BKR", 2500.00, 2950.00, 18.0, "85072000", "PCS", "", 2, 5, 10, "Active"],
        ["AC12345", "BKR 1800W High Pressure Washer 140 Bar", "Power Tools", "BKR", 7203.39, 8500.00, 18.0, "84243000", "PCS", "8901234567891", 10, 20, 50, "Active"],
        ["HM0486", "GORILLA HEAVY DUTY TAPE 6035060", "Hardware", "Gorilla", 1134.75, 1339.00, 18.0, "39191000", "PCS", "", 20, 40, 100, "Active"],
        ["CA0147", "BKR Portable Hydraulic 2 Ton Car Jack", "Car Accessories", "BKR", 5084.75, 6000.00, 18.0, "84253900", "PCS", "", 5, 15, 30, "Active"],
        ["TL0912", "Cordless Brushless Drill 21V Dual Battery", "Power Tools", "BKR", "", 4500.00, 18.0, "84672100", "SET", "", 5, 10, 15, "Active"],
    ]
    
    for row in sample_rows:
        ws_data.append(row)
        
    # Styling for Sheet 1
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    for col_idx in range(1, len(header_titles) + 1):
        cell = ws_data.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        ws_data.row_dimensions[1].height = 28
        
    for row_idx in range(2, len(sample_rows) + 2):
        ws_data.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(header_titles) + 1):
            cell = ws_data.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border
            
            # Alignments & Formats
            if col_idx in [5, 6]:  # Rate, Price
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [7, 11, 12, 13]:  # GST, Min, Reorder, Initial Stock
                cell.number_format = "#,##0.00" if col_idx == 7 else "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [1, 8, 9, 10, 14]:  # SKU, HSN, Unit, Barcode, Status
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Auto-fit column widths
    for col in ws_data.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    # -------------------------------------------------------------
    # Sheet 2: Field Guide & Rules
    # -------------------------------------------------------------
    ws_guide = wb.create_sheet(title="Field Guide & Rules")
    ws_guide.views.sheetView[0].showGridLines = True
    
    guide_headers = ["Field Name", "Required?", "Accepted Column Aliases", "Data Type", "Default Value", "Description & Rules"]
    ws_guide.append(guide_headers)
    
    guide_rows = [
        ["SKU", "YES", "SKU, Product SKU, Item SKU, Item Code, Part Number", "Text", "-", "Unique product code. Matches pattern: 2 letters + 4+ digits (e.g. LG0831). Variants with suffixes like LG0831-1 are strictly isolated as separate products."],
        ["Product Name", "YES", "Product Name, Name, Item Name, Title, Description", "Text", "-", "Full descriptive name of the product."],
        ["Category", "NO", "Category, Product Category, Item Category, Group", "Text", "Empty", "Classification category (e.g. Power Tools, Garden Equipment, Spares)."],
        ["Brand", "NO", "Brand, Make, Manufacturer", "Text", "Empty", "Brand name (e.g. BKR, JSPL, Gardena, Bosch)."],
        ["Rate (Excl GST)", "NO", "Rate (Excl GST), Item Rate, Rate, Selling Price, Basic Rate", "Decimal", "0.00", "Base unit selling price excluding tax. If left blank, automatically deduced from Price (Incl GST)."],
        ["Price (Incl GST)", "NO", "Price (Incl GST), Price Incl GST, MRP, Gross Price, Gross Rate", "Decimal", "0.00", "Total selling price including GST. If Rate is blank, formula applied: Rate = Price / (1 + GST% / 100)."],
        ["GST Rate (%)", "NO", "GST Rate, GST %, GST, Tax Rate, Tax %", "Decimal", "18.0", "Applicable GST percentage (e.g. 0, 5, 12, 18, 28)."],
        ["HSN Code", "NO", "HSN, HSN Code, HSN/SAC, SAC", "Text", "Empty", "Standard 4 to 8 digit HSN/SAC code."],
        ["Unit", "NO", "Unit, UOM, Unit of Measure", "Text", "PCS", "Measurement unit (e.g. PCS, SET, NOS, KG, MTR, BOX)."],
        ["Barcode", "NO", "Barcode, EAN, UPC", "Text", "Empty", "Unique product barcode or EAN number."],
        ["Min Stock Level", "NO", "Min Stock Level, Min Stock, Safety Stock", "Integer", "0", "Minimum inventory threshold before triggering low-stock warnings."],
        ["Reorder Level", "NO", "Reorder Level, Reorder Point, ROP", "Integer", "0", "Suggested reorder quantity threshold."],
        ["Initial Stock", "NO", "Initial Stock, Opening Stock, Quantity, Qty, Stock", "Integer", "0", "Initial quantity to credit into the company's Central Warehouse."],
        ["Status", "NO", "Status, Product Status", "Text", "Active", "Product status: 'Active' or 'Inactive'."]
    ]
    
    for r in guide_rows:
        ws_guide.append(r)
        
    for col_idx in range(1, len(guide_headers) + 1):
        cell = ws_guide.cell(row=1, column=col_idx)
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
        ws_guide.row_dimensions[1].height = 26
        
    for row_idx in range(2, len(guide_rows) + 2):
        ws_guide.row_dimensions[row_idx].height = 24
        for col_idx in range(1, len(guide_headers) + 1):
            cell = ws_guide.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = cell_border
            if col_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Calibri", size=10, bold=True, color="DC2626" if cell.value == "YES" else "4B5563")
            elif col_idx in [1, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
    ws_guide.column_dimensions["A"].width = 20
    ws_guide.column_dimensions["B"].width = 14
    ws_guide.column_dimensions["C"].width = 30
    ws_guide.column_dimensions["D"].width = 14
    ws_guide.column_dimensions["E"].width = 14
    ws_guide.column_dimensions["F"].width = 50
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Excel template created successfully at: {output_path}")

if __name__ == "__main__":
    create_product_import_template("backend/product_import_template.xlsx")
    create_product_import_template("backend/templates/product_import_template.xlsx")
